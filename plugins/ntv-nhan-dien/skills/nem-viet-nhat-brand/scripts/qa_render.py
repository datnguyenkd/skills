# -*- coding: utf-8 -*-
"""QA: render file .pptx ra ảnh PNG từng slide để soi lỗi tràn chữ / đè / lệch.

Chạy:  python3 qa_render.py <file.pptx> [thư-mục-ra] [--dpi 110] [--quicklook]
Mặc định ảnh ra ở  <thư-mục-của-file>/_qa/slide-01.png, slide-02.png, ...

Cách render theo hệ (tự chọn):
- macOS   : (1) LibreOffice headless → PDF → pdftoppm → PNG   nếu máy có LibreOffice + poppler
            (2) DỰ PHÒNG KHÔNG CẦN CÀI GÌ: tách deck thành từng file 1-slide rồi dùng Quick Look
                (qlmanage -t) của macOS render từng slide. Đủ để soi tràn/đè; font per-user nhận đúng.
- Linux   : LibreOffice headless + pdftoppm.
- Windows : PowerPoint qua COM (PowerShell), không cần LibreOffice.
Nếu thiếu công cụ, script báo rõ lệnh cài rồi thoát (không đoán mò).
"""
import os, sys, glob, shutil, platform, subprocess, tempfile

SOFFICE_CANDIDATES = [
    "soffice", "libreoffice",
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    os.path.expanduser("~/Applications/LibreOffice.app/Contents/MacOS/soffice"),
    "/opt/homebrew/bin/soffice", "/usr/local/bin/soffice",
    "/usr/bin/soffice", "/usr/lib/libreoffice/program/soffice",
]


def find_soffice():
    for c in SOFFICE_CANDIDATES:
        p = shutil.which(c) if os.sep not in c else (c if os.path.exists(c) else None)
        if p:
            return p
    return None


def _clear_dir(d):
    os.makedirs(d, exist_ok=True)
    for f in glob.glob(os.path.join(d, "slide-*.png")):
        os.remove(f)


# ---------- macOS / Linux: LibreOffice ----------
def render_libreoffice(pptx, out_dir, dpi, soffice, pdftoppm):
    _clear_dir(out_dir)
    with tempfile.TemporaryDirectory() as tmp:
        # profile riêng để không đụng LibreOffice đang mở của người dùng
        profile = "file://" + os.path.join(tmp, "lo-profile")
        cmd = [soffice, f"-env:UserInstallation={profile}", "--headless",
               "--convert-to", "pdf", "--outdir", tmp, os.path.abspath(pptx)]
        r = subprocess.run(cmd, capture_output=True, text=True, timeout=300)
        pdfs = glob.glob(os.path.join(tmp, "*.pdf"))
        if not pdfs:
            print("LibreOffice không ra PDF.\nSTDOUT:", r.stdout, "\nSTDERR:", r.stderr)
            sys.exit(3)
        subprocess.run([pdftoppm, "-png", "-r", str(dpi), pdfs[0], os.path.join(out_dir, "slide")], check=True)
    # chuẩn hoá tên slide-1.png → slide-01.png
    for f in sorted(glob.glob(os.path.join(out_dir, "slide-*.png"))):
        num = os.path.basename(f)[6:-4]
        if num.isdigit() and len(num) < 2:
            os.rename(f, os.path.join(out_dir, f"slide-{int(num):02d}.png"))


# ---------- macOS dự phòng: Quick Look từng slide ----------
def render_quicklook(pptx, out_dir, dpi):
    from pptx import Presentation
    if not shutil.which("qlmanage"):
        print("Không thấy qlmanage (Quick Look) — chỉ có trên macOS."); sys.exit(2)
    _clear_dir(out_dir)
    width = int(13.333 * dpi)
    n = len(Presentation(pptx).slides)
    with tempfile.TemporaryDirectory() as tmp:
        for i in range(n):
            prs = Presentation(pptx)
            lst = prs.slides._sldIdLst
            for j, sid in enumerate(list(lst)):
                if j != i:
                    prs.part.drop_rel(sid.rId); lst.remove(sid)
            one = os.path.join(tmp, f"s{i+1:02d}.pptx"); prs.save(one)
            subprocess.run(["qlmanage", "-t", "-s", str(width), "-o", tmp, one],
                           capture_output=True, timeout=120)
            png = one + ".png"
            if os.path.exists(png):
                shutil.move(png, os.path.join(out_dir, f"slide-{i+1:02d}.png"))
            else:
                print(f"  ! Quick Look không render được slide {i+1}")


# ---------- Windows: PowerPoint COM ----------
def render_windows(pptx, out_dir, dpi):
    _clear_dir(out_dir)
    w = int(13.333 * dpi); h = int(7.5 * dpi)
    ps = f'''
$pp=New-Object -ComObject PowerPoint.Application
$d=$pp.Presentations.Open("{os.path.abspath(pptx)}",$true,$false,$false)
$i=1; foreach($sl in $d.Slides){{ $sl.Export("{os.path.abspath(out_dir)}\\slide-" + $i.ToString("00") + ".png","PNG",{w},{h}); $i++ }}
$d.Close(); $pp.Quit()
'''
    subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True)


def main():
    argv = sys.argv[1:]
    dpi = 110
    if "--dpi" in argv:
        k = argv.index("--dpi"); dpi = int(argv[k + 1]); del argv[k:k + 2]
    force_ql = "--quicklook" in argv
    args = [a for a in argv if not a.startswith("--")]
    if not args:
        print(__doc__); sys.exit(1)
    pptx = args[0]
    if not os.path.exists(pptx):
        print("Không thấy file:", pptx); sys.exit(1)
    out_dir = args[1] if len(args) > 1 else os.path.join(os.path.dirname(os.path.abspath(pptx)), "_qa")

    sysname = platform.system()
    if sysname == "Windows":
        render_windows(pptx, out_dir, dpi); how = "PowerPoint COM"
    else:
        soffice, pdftoppm = find_soffice(), shutil.which("pdftoppm")
        if soffice and pdftoppm and not force_ql:
            render_libreoffice(pptx, out_dir, dpi, soffice, pdftoppm); how = "LibreOffice + pdftoppm"
        elif sysname == "Darwin":
            if not force_ql:
                print("Không có LibreOffice/poppler → dùng Quick Look (macOS) render từng slide.")
            render_quicklook(pptx, out_dir, dpi); how = "Quick Look (qlmanage)"
        else:
            print("THIẾU công cụ render, chưa QA được:")
            if not soffice:  print("  - LibreOffice  →  apt/dnf install libreoffice")
            if not pdftoppm: print("  - poppler (pdftoppm)  →  apt/dnf install poppler-utils")
            sys.exit(2)

    pngs = sorted(glob.glob(os.path.join(out_dir, "slide-*.png")))
    print(f"Đã render {len(pngs)} slide bằng {how} → {out_dir}")
    for p in pngs:
        print("  ", p)


if __name__ == "__main__":
    main()
